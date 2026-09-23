#!/usr/bin/env python3
"""Export the 'Final data KZOO' workbooks into the tidy CSV files the website reads.

The workbooks stay the source of truth. This script only reads them; it never
writes to them. Run it after each quarterly update, then commit the CSVs.

    python tools/export_from_excel.py --source "D:/Downtown/Final data KZOO"

Writes into ./data :
    traffic_monthly.csv    corridor x direction x month
    people_monthly.csv     segment x month
    parking_monthly.csv    facility x time window x month
    transit_monthly.csv    stop x month

Requires openpyxl:  pip install openpyxl
"""
import argparse
import csv
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required.  Install it with:  pip install openpyxl")

# --------------------------------------------------------------------------- corrections
# Documented departures from the raw cells. Each one is listed on the site's
# Data page and in data/CORRECTIONS.md so the numbers can be traced back.
EMPLOYEE_FIX = {("2025-07"): 15598}     # Employees!H24 reads 2; the workbook's own copy at H56 has 15,598
CORRIDOR_RENAME = {"Bus US-31": "Bus US-131"}   # Outgoing sheet spells this corridor differently

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
FILES = {
    "traffic": "Traffic volumes.xlsx",
    "people": "No. of visitors , employees , residents.xlsx",
    "parking": "Parking_Diana.xlsx",
    "transit": "Transit data DA.xlsx",
}
PARKING_SHEETS = {"Afternoon": "Weekday afternoon",
                  "weekday evenings": "Weekday evening",
                  "weekend evenings": "Weekend evening"}
PARKING_BLOCKS = {"On-street": (5, 7, "{}"), "Ramps": (11, 12, "Ramp {}"), "Surface lots": (16, 21, "Lot {}")}
TRANSIT_SHEETS = {2023: "2023 data", 2024: "2024 Data", 2025: "2025 Data", 2026: "2026 Data"}

skipped = []


def num(cell_value, where=""):
    """Return a float, or None for blanks and anything non-numeric (logged)."""
    if cell_value is None or cell_value == "":
        return None
    if isinstance(cell_value, (int, float)):
        return float(cell_value)
    skipped.append(f"{where}: {cell_value!r}")
    return None


def ym(year, month):
    return f"{year}-{month:02d}"


def book(source, key):
    path = os.path.join(source, FILES[key])
    if not os.path.exists(path):
        sys.exit(f"Could not find {path}")
    return openpyxl.load_workbook(path, data_only=True)


# --------------------------------------------------------------------------- traffic
def export_traffic(source, out):
    wb = book(source, "traffic")
    rows = []
    for sheet, direction in (("Incoming traffic", "incoming"), ("Outgoing traffic", "outgoing")):
        ws = wb[sheet]
        for r in range(1, ws.max_row + 1):
            if str(ws.cell(r, 1).value).strip() != "Zone Name":
                continue                                   # each year starts with a 'Zone Name' header row
            year = int(ws.cell(r - 1, 2).value)
            rr = r + 1
            while ws.cell(rr, 1).value not in (None, ""):
                raw = str(ws.cell(rr, 1).value).strip()
                corridor = CORRIDOR_RENAME.get(raw, raw)
                for m in range(1, 13):
                    v = num(ws.cell(rr, 1 + m).value, f"{sheet} {corridor} {ym(year, m)}")
                    if v is not None:
                        rows.append({"month": ym(year, m), "corridor": corridor,
                                     "direction": direction, "vehicles_per_day": int(round(v))})
                rr += 1
    rows.sort(key=lambda d: (d["month"], d["corridor"], d["direction"]))
    write(out, "traffic_monthly.csv", ["month", "corridor", "direction", "vehicles_per_day"], rows)


# --------------------------------------------------------------------------- people
def export_people(source, out):
    wb = book(source, "people")
    field = {"Visits": "visits", "Visit Frequency": "visits_per_person",
             "Avg. Dwell Time": "avg_dwell_minutes", "Panel Visits": "panel_visits",
             "Visits YoY": "visits_yoy"}
    out_rows = {}
    for ws in wb.worksheets:
        segment = ws.title.lower()                          # residents / visitors / employees
        for r in range(1, 41):                              # rows 44+ repeat values for the workbook's charts
            if str(ws.cell(r, 1).value).strip() != "Metrics":
                continue
            year = int(re.search(r"(20\d\d)", str(ws.cell(r - 1, 1).value)).group(1))
            rr = r + 1
            while rr <= 40 and ws.cell(rr, 1).value not in (None, "") and not re.search(r"20\d\d", str(ws.cell(rr, 1).value)):
                metric = str(ws.cell(rr, 1).value).strip()
                key = field.get(metric)
                if metric.lower() == ws.title.lower():      # 'Visitors' / 'Employees' / 'Residents' = the people count
                    key = "people"
                if key:
                    for m in range(1, 13):
                        v = num(ws.cell(rr, 1 + m).value, f"{ws.title} {metric} {ym(year, m)}")
                        if v is None:
                            continue
                        row = out_rows.setdefault((ym(year, m), segment),
                                                  {"month": ym(year, m), "segment": segment})
                        row[key] = v
                rr += 1
    # documented correction
    fix = out_rows.get(("2025-07", "employees"))
    if fix and fix.get("people") != EMPLOYEE_FIX["2025-07"]:
        fix["people"] = EMPLOYEE_FIX["2025-07"]
        fix["corrected"] = "people"
    rows = []
    for row in out_rows.values():
        rows.append({
            "month": row["month"], "segment": row["segment"],
            "visits": fmt(row.get("visits"), 0), "people": fmt(row.get("people"), 0),
            "visits_per_person": fmt(row.get("visits_per_person"), 2),
            "avg_dwell_minutes": fmt(row.get("avg_dwell_minutes"), 0),
            "panel_visits": fmt(row.get("panel_visits"), 0),
            "visits_yoy": fmt(row.get("visits_yoy"), 4),
            "corrected": row.get("corrected", ""),
        })
    rows.sort(key=lambda d: (d["month"], d["segment"]))
    write(out, "people_monthly.csv",
          ["month", "segment", "visits", "people", "visits_per_person", "avg_dwell_minutes",
           "panel_visits", "visits_yoy", "corrected"], rows)


# --------------------------------------------------------------------------- parking
def export_parking(source, out):
    wb = book(source, "parking")
    rows = []
    for sheet, window in PARKING_SHEETS.items():
        ws = wb[sheet]
        for ftype, (r0, r1, label) in PARKING_BLOCKS.items():
            for r in range(r0, r1 + 1):
                facility = label.format(str(ws.cell(r, 1).value).strip())
                spaces = num(ws.cell(r, 2).value, f"{sheet} {facility} inventory")
                if spaces is None:
                    continue
                for i in range(48):                         # column C = Jan 2023, 48 months to Dec 2026
                    year, month = 2023 + i // 12, i % 12 + 1
                    v = num(ws.cell(r, 3 + i).value, f"{sheet} {facility} {ym(year, month)}")
                    if v is None:
                        continue                            # blank month, or a stray non-numeric cell
                    rows.append({"month": ym(year, month), "time_window": window,
                                 "facility_type": ftype, "facility": facility,
                                 "spaces": int(spaces), "occupied": int(round(v))})
    rows.sort(key=lambda d: (d["month"], d["time_window"], d["facility_type"], d["facility"]))
    write(out, "parking_monthly.csv",
          ["month", "time_window", "facility_type", "facility", "spaces", "occupied"], rows)


# --------------------------------------------------------------------------- transit
def export_transit(source, out):
    wb = book(source, "transit")
    names, data = {}, {}
    for year, sheet in TRANSIT_SHEETS.items():
        ws = wb[sheet]
        for kind, r0 in (("alightings", 3), ("boardings", 60)):
            for r in range(r0, r0 + 54):
                stop = ws.cell(r, 1).value
                if stop is None:
                    continue
                stop = int(stop)
                names.setdefault(stop, str(ws.cell(r, 2).value).strip())
                for m in range(1, 13):
                    v = num(ws.cell(r, 2 + m).value, f"{sheet} stop {stop} {ym(year, m)}")
                    if v is None:
                        continue                            # blank or '-' = no service / not reported
                    data.setdefault((ym(year, m), stop), {})[kind] = int(round(v))
    rows = [{"month": k[0], "stop_id": k[1], "stop_name": names.get(k[1], ""),
             "boardings": v.get("boardings", ""), "alightings": v.get("alightings", "")}
            for k, v in data.items()]
    rows.sort(key=lambda d: (d["month"], d["stop_id"]))
    write(out, "transit_monthly.csv", ["month", "stop_id", "stop_name", "boardings", "alightings"], rows)


# --------------------------------------------------------------------------- helpers
def fmt(v, places):
    if v is None:
        return ""
    return str(int(round(v))) if places == 0 else f"{v:.{places}f}".rstrip("0").rstrip(".")


def write(out, name, header, rows):
    os.makedirs(out, exist_ok=True)
    path = os.path.join(out, name)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=header)
        w.writeheader()
        w.writerows(rows)
    print(f"  {name:<24} {len(rows):>6} rows")


def main():
    ap = argparse.ArgumentParser(description="Export the Final data KZOO workbooks to the website's CSV files.")
    ap.add_argument("--source", default=r"D:\Downtown\Final data KZOO", help="folder holding the four workbooks")
    ap.add_argument("--out", default="data", help="folder to write the CSV files into")
    a = ap.parse_args()
    print(f"Reading {a.source}")
    export_traffic(a.source, a.out)
    export_people(a.source, a.out)
    export_parking(a.source, a.out)
    export_transit(a.source, a.out)
    if skipped:
        print(f"\n{len(skipped)} cell(s) skipped as non-numeric:")
        for s in skipped[:20]:
            print("   ", s)
    print("\nDone. Review the CSVs, then commit them.")


if __name__ == "__main__":
    main()
